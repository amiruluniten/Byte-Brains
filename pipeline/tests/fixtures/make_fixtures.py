"""Generate the small raw xlsx fixtures for offline tests.

Miniature versions of the real DOSM TSA workbooks, mirroring the layout documented
in research/tsa-xlsx-map.md: bilingual merged A:B label cells, year header in row 3,
multi-block sheets (values / % change / % contribution), revision flags in year
headers ("2023p"), footnotes ("n.a", "4.1*"), phantom max_col artifacts.

Run from the pipeline directory:
    .venv/bin/python tests/fixtures/make_fixtures.py
"""
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent


def add_phantom_column(ws):
    """Reproduce the 2024-file artifact: formatting reaches far beyond the data."""
    ws.column_dimensions["XFD"].width = 8.43


def indicator_inbound_2024(wb):
    """Visitor basis (2019-2024): A1 visitors, A2 tourists, A3 excursionists."""
    ws = wb.create_sheet("Indicator Inbound")
    ws["A3"] = "Tahun\n Year"
    for col, year in zip("CDEFGH", ["2019", 2020, 2021, 2022, "2023", "2024"]):
        ws[f"{col}3"] = year
    ws.merge_cells("A5:B5")
    ws["A5"] = "A. Ketibaan Pelawat\n    Visitor arrivals"
    ws["D5"] = "Bilangan orang\nNumber of persons"

    ws["B7"] = "A1. Ketibaan pelawat ke Malaysia dari negara terpilih\n      Visitor arrivals to Malaysia from selected countries"
    for col, v in zip("CDEFGH", [35045625, 6101378, 399865, 14267416, 28964308, 37961485]):
        ws[f"{col}7"] = v
    ws["B9"] = "Singapore"
    for col, v in zip("CDEFGH", [17033066, 2871340, 16729, 8399088, 16496436, 23046245]):
        ws[f"{col}9"] = v

    ws["B26"] = "A2. Ketibaan pelancong ke Malaysia\n      Tourist arrivals to Malaysia"
    for col, v in zip("CDEFGH", [26100784, 4332722, 134728, 10070964, 20141846, 25016698]):
        ws[f"{col}26"] = v
    ws["B27"] = "A3. Ketibaan pelawat harian ke Malaysia\n      Excursionist arrivals to Malaysia"
    for col, v in zip("CDEFGH", [8944841, 1768656, 265137, 4196452, 8822462, 12944787]):
        ws[f"{col}27"] = v

    ws["A42"] = "Purata bilangan hari menginap / Average length of stay (ALOS)"
    ws["C42"] = 7.4
    ws["D42"] = "4.1*"  # asterisk footnote
    ws["E42"] = "n.a"  # not available
    add_phantom_column(ws)


def jad_1a_2024(wb):
    """Inbound tourism consumption (tourists) by product, RM million, 2015-2024."""
    ws = wb.create_sheet("Jad 1A")
    for col, year in zip("BCDEFGHIJK", ["2015", "2016", "2017", "2018", "2019", 2020, 2021, 2022, "2023", "2024"]):
        ws[f"{col}3"] = year
    ws["A5"] = "Produk\n  Products"
    ws["B5"] = "RM Juta\nRM Million"
    ws["A6"] = "Perkhidmatan penginapan\nAccommodation services"
    for col, v in zip("BCDEFGHIJK", [17656.4, 20142.5, 21034.2, 21622.8, 22007.3, 3144, 64.8, 5234.1, 14369, 19752.7]):
        ws[f"{col}6"] = v
    ws["A14"] = "Jumlah\nTotal"
    for col, v in zip("BCDEFGHIJK", [72592.5, 79325.9, 82921.5, 84929.9, 86706.5, 13157.3, 389.8, 32473.3, 72992.8, 102815.3]):
        ws[f"{col}14"] = v
    ws["A15"] = "Perubahan peratusan tahunan (%)\nAnnual percentage change (%)"
    ws["B15"] = ".."  # first year of % change: not applicable
    ws["C15"] = 9.3
    ws.merge_cells("B17:L17")
    ws["B17"] = "Peratus sumbangan (%)\nPercentage share (%)"
    ws["A26"] = "Jumlah\nTotal"  # second block total (share block)
    for col in "BCDEFGHIJK":
        ws[f"{col}26"] = 100
    add_phantom_column(ws)


def indicator_inbound_2023(wb):
    """Tourist basis (2015-2023): A1 tourist arrivals by country, continent-grouped."""
    ws = wb.create_sheet("Indicator Inbound")
    ws["B3"] = "Tahun\n Year"
    for col, year in zip("CDEFGHIJK", ["2015", "2016", "2017", "2018", "2019", "2020", 2021, "2022", "2023"]):
        ws[f"{col}3"] = year
    ws.merge_cells("A5:B5")
    ws["A5"] = "A. Ketibaan pelancong\n     Tourist arrivals"
    ws["C5"] = "Bilangan orang\nNo of persons"

    ws["B7"] = "A1. Ketibaan pelancong ke Malaysia dari negara terpilih\n      Tourist arrivals to Malaysia from selected countries"
    for col, v in zip("CDEFG", [25721251, 26757392, 25948459, 25832354, 26100784]):
        ws[f"{col}7"] = v
    for col, v in zip("HIJK", [4332722, 134728, 10070964, 20141846]):
        ws[f"{col}7"] = v
    ws["B8"] = "i. Benua Asia / Asia Continent"  # continent group header, no values
    ws["B9"] = "Singapore"
    for col, v in zip("CDEF", [12930754, 13272961, 12441713, 12010411]):
        ws[f"{col}9"] = v


def table_1a_2023(wb):
    """Same as Jad 1A but 2015-2023 with the "2023p" preliminary year header."""
    ws = wb.create_sheet("table 1A")
    for col, year in zip("BCDEFGHIJ", ["2015", "2016", "2017", "2018", "2019", 2020, 2021, 2022, "2023p"]):
        ws[f"{col}3"] = year
    ws["A5"] = "Produk\n  Products"
    ws["B5"] = "RM Juta\nRM Million"
    ws["A14"] = "Jumlah\nTotal"
    for col, v in zip("BCDEFGHIJ", [72592.5, 79325.9, 82921.5, 84929.9, 86706.5, 13157.3, 389.8, 32473.3, 72992.8]):
        ws[f"{col}14"] = v
    ws["A15"] = "Perubahan peratusan tahunan (%)\nAnnual percentage change (%)"
    ws["B15"] = ".."
    ws["C15"] = 9.3


wb24 = Workbook()
wb24.remove(wb24.active)
indicator_inbound_2024(wb24)
jad_1a_2024(wb24)
wb24.save(HERE / "tourism_2024.fixture.xlsx")

wb23 = Workbook()
wb23.remove(wb23.active)
indicator_inbound_2023(wb23)
table_1a_2023(wb23)
wb23.save(HERE / "tourism_2023.fixture.xlsx")

print("fixtures written to", HERE)
