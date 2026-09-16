"""Slice 2 (red): the TSA national-series extractor, run offline against fixtures.

Fixtures are miniature versions of the real workbooks (see
tests/fixtures/make_fixtures.py and research/tsa-xlsx-map.md). Ground-truth values
are the ticket's acceptance numbers.
"""
from openpyxl import load_workbook

from bytebrains_pipeline.extractors import tsa_national


def load(fixtures_dir, name):
    return load_workbook(fixtures_dir / name, data_only=True)


class TestIndicatorInbound2024:
    """Visitor basis (2019-2024): A1 visitors, A2 tourists, A3 excursionists."""

    def test_visitor_arrivals_total(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_visitor_arrivals(wb)
        assert s.basis == "visitor"
        assert s.unit == "persons"
        assert s.window == "2019-2024"
        assert [(o.year, o.value) for o in s.values] == [
            (2019, 35045625),
            (2020, 6101378),
            (2021, 399865),
            (2022, 14267416),
            (2023, 28964308),
            (2024, 37961485),
        ]

    def test_tourist_arrivals_total(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_tourist_arrivals_2024(wb)
        assert s.basis == "tourist"
        assert s.window == "2019-2024"
        assert s.values[0].value == 26100784  # 2019 ground truth

    def test_excursionist_arrivals_total(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_excursionist_arrivals(wb)
        assert s.basis == "excursionist"
        assert s.window == "2019-2024"
        assert s.values[0].value == 8944841


class TestIndicatorInbound2023:
    """Tourist basis (2015-2023): A1 tourist arrivals total."""

    def test_tourist_arrivals_total(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2023.fixture.xlsx")
        s = tsa_national.extract_tourist_arrivals_2023(wb)
        assert s.basis == "tourist"
        assert s.window == "2015-2023"
        assert [(o.year, o.value) for o in s.values] == [
            (2015, 25721251),
            (2016, 26757392),
            (2017, 25948459),
            (2018, 25832354),
            (2019, 26100784),
            (2020, 4332722),
            (2021, 134728),
            (2022, 10070964),
            (2023, 20141846),
        ]


class TestJad1A:
    """Inbound tourism consumption (tourists) by product, RM million."""

    def test_totals_2015_2024(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_inbound_consumption(wb)
        assert s.basis == "tourist"
        assert s.unit == "rm_million"
        assert s.window == "2015-2024"
        assert [(o.year, o.value) for o in s.values] == [
            (2015, 72592.5),
            (2016, 79325.9),
            (2017, 82921.5),
            (2018, 84929.9),
            (2019, 86706.5),
            (2020, 13157.3),
            (2021, 389.8),
            (2022, 32473.3),
            (2023, 72992.8),
            (2024, 102815.3),
        ]

    def test_reads_values_block_not_share_block(self, fixtures_dir):
        # Multi-block sheet: the second "Jumlah/Total" row (row 26, share block) is 100s.
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_inbound_consumption(wb)
        assert all(o.value != 100 for o in s.values)

    def test_revision_flag_parsed_from_year_header(self, fixtures_dir):
        # 2023 file marks the last year "2023p" (preliminary).
        wb = load(fixtures_dir, "tourism_2023.fixture.xlsx")
        s = tsa_national.extract_inbound_consumption(wb)
        assert s.values[-1].year == 2023
        assert s.values[-1].revision_flag == "p"
        assert s.values[-1].value == 72992.8
        assert all(o.revision_flag is None for o in s.values[:-1])


class TestTsa2025Edition:
    """Ticket #13: the TSA 2025 workbook (2025p preliminary release).

    Same anchors and year-header row as the 2024 edition; blank spacer rows
    between sections are tolerated. The 2025 observations are flagged
    preliminary and the 2024 consumption observation carries the workbook's
    revision (RM102,931.3m vs the 2024 edition's RM102,815.3m).
    """

    def test_visitor_arrivals_2025_preliminary(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        s = tsa_national.extract_visitor_arrivals_2025(wb)
        assert s.series_id == "arrivals_visitor_2019_2025"
        assert s.basis == "visitor"
        assert s.unit == "persons"
        assert s.window == "2019-2025"
        assert s.source.file == "tourism_2025.xlsx"
        assert [(o.year, o.value) for o in s.values] == [
            (2019, 35045625),
            (2020, 6101378),
            (2021, 399865),
            (2022, 14267416),
            (2023, 28964308),
            (2024, 37961485),
            (2025, 42196892),
        ]
        last = s.values[-1]
        assert last.revision_status == "preliminary"
        assert all(o.revision_status == "final" for o in s.values[:-1])

    def test_tourist_arrivals_2025_preliminary(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        s = tsa_national.extract_tourist_arrivals_2025(wb)
        assert s.series_id == "arrivals_tourist_2019_2025"
        assert s.basis == "tourist"
        assert s.window == "2019-2025"
        assert s.values[-1].value == 26613597
        assert s.values[-1].revision_status == "preliminary"

    def test_excursionist_arrivals_2025_preliminary(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        s = tsa_national.extract_excursionist_arrivals_2025(wb)
        assert s.series_id == "arrivals_excursionist_2019_2025"
        assert s.basis == "excursionist"
        assert s.window == "2019-2025"
        assert s.values[-1].value == 15583295
        assert s.values[-1].revision_status == "preliminary"

    def test_inbound_consumption_2015_2025_with_revised_2024(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        s = tsa_national.extract_inbound_consumption_2025(wb)
        assert s.series_id == "inbound_consumption_tourist_2015_2025"
        assert s.basis == "tourist"
        assert s.unit == "rm_million"
        assert s.window == "2015-2025"
        assert s.source.file == "tourism_2025.xlsx"
        by_year = {o.year: o for o in s.values}
        assert by_year[2024].value == 102931.3  # revised (2024 edition: 102815.3)
        assert by_year[2024].revision_status == "revised"
        assert by_year[2024].revision_flag is None  # the 2025 workbook header is plain "2024"
        assert by_year[2025].value == 119312.0
        assert by_year[2025].revision_status == "preliminary"
        assert by_year[2025].revision_flag == "p"  # the workbook's own "2025p" header
        assert by_year[2019].value == 86706.5
        assert by_year[2019].revision_status == "final"

    def test_reads_values_block_not_share_block(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        s = tsa_national.extract_inbound_consumption_2025(wb)
        assert all(o.value != 100 for o in s.values)

    def test_blank_spacer_rows_tolerated(self, fixtures_dir):
        # The 2025 export inserts blank spacer rows between sections (rows 8,
        # 25, 28-29); the anchor/year-header strategy must skip them.
        wb = load(fixtures_dir, "tourism_2025.fixture.xlsx")
        ws = wb["Indicator Inbound"]
        assert ws["B8"].value is None and ws["B25"].value is None
        s = tsa_national.extract_visitor_arrivals_2025(wb)
        assert len(s.values) == 7  # 2019-2025, spacers skipped


class TestGotchaHelpers:
    def test_footnote_strings_become_null(self, fixtures_dir):
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        ws = wb["Indicator Inbound"]
        assert tsa_national.to_number(ws["D42"].value) == 4.1  # "4.1*" keeps the number
        assert tsa_national.to_number(ws["E42"].value) is None  # "n.a"

    def test_phantom_columns_truncated(self, fixtures_dir):
        # Real 2024 file reports max_col=16384; extractor must ignore it.
        wb = load(fixtures_dir, "tourism_2024.fixture.xlsx")
        s = tsa_national.extract_visitor_arrivals(wb)
        assert len(s.values) == 6
