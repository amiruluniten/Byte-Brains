"""Bundle assembly and emission."""
from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path

import subprocess

from openpyxl import load_workbook

from .bundle import (
    BUNDLE_VERSION,
    Bundle,
    MacroSeriesFragment,
    NationalSeriesFragment,
)
from .extractors import cpi, inbrief_markets, tsa_national
from .missing_billions import LAND_MODE_SHARE_2024_PCT, compute_missing_billions

# ticket T5: default location of the WEF TTDI indicator file (checked into the repo)
_REPO_ROOT = Path(__file__).resolve().parents[3]
_DEFAULT_WEF_CSV = _REPO_ROOT / "source" / "dataset" / "WEF_TTDI.csv"


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _inbrief_text(data_dir: Path) -> Path:
    """The pdftotext -layout extract of the In Brief PDF; generated on the fly if absent."""
    txt = data_dir / "inbrief2024.txt"
    if not txt.exists():
        pdf = data_dir / "inbrief2024.pdf"
        if not pdf.exists():
            raise FileNotFoundError(
                f"neither {txt} nor {pdf} exists; run: pdftotext -layout inbrief2024.pdf inbrief2024.txt"
            )
        subprocess.run(
            ["pdftotext", "-layout", str(pdf), str(txt)],
            check=True,
        )
    return txt


def build_bundle(data_dir: Path, wef_csv: Path | None = None) -> Bundle:
    """Extract the national series from the raw TSA workbooks and build the bundle."""
    file_2023 = data_dir / "tourism_2023.xlsx"
    file_2024 = data_dir / "tourism_2024.xlsx"
    file_2025 = data_dir / "tourism_2025.xlsx"  # ticket #13: TSA 2025 ("2025p")
    file_inbrief = _inbrief_text(data_dir)
    wb23 = load_workbook(file_2023, data_only=True)
    wb24 = load_workbook(file_2024, data_only=True)
    wb25 = load_workbook(file_2025, data_only=True) if file_2025.exists() else None

    fragment = NationalSeriesFragment(
        series=[
            tsa_national.extract_tourist_arrivals_2023(wb23),
            tsa_national.extract_visitor_arrivals(wb24),
            tsa_national.extract_tourist_arrivals_2024(wb24),
            tsa_national.extract_excursionist_arrivals(wb24),
            tsa_national.extract_inbound_consumption(wb24),
        ]
    )
    # ticket #13: the TSA 2025 edition enters ADDITIVELY — new series with their
    # own windows, existing series ids/windows unchanged. Later official
    # workbook wins for the counterfactual (its 2024 receipts are revised).
    if wb25 is not None:
        fragment.series.extend(
            [
                tsa_national.extract_visitor_arrivals_2025(wb25),
                tsa_national.extract_tourist_arrivals_2025(wb25),
                tsa_national.extract_excursionist_arrivals_2025(wb25),
                tsa_national.extract_inbound_consumption_2025(wb25),
            ]
        )

    # cross-file consistency: 2023 edition table 1A vs 2024 edition Jad 1A
    legacy = tsa_national.extract_inbound_consumption(wb23)

    # ticket T3: source-market panel from the Tourism Malaysia In Brief 2024
    source_market = inbrief_markets.build_source_market_fragment(
        file_inbrief.read_text(), source_file=file_inbrief.name
    )

    # ticket T4: national CPI deflator + the Missing Billions counterfactual
    file_cpi = data_dir / "cpi_headline.csv"
    if not file_cpi.exists():
        raise FileNotFoundError(
            f"{file_cpi} missing; re-fetch with: dosm download cpi_headline"
        )
    cpi_series = cpi.extract_cpi(file_cpi)
    macro_series_series = [cpi_series]
    # ticket #13: when the 2025 workbook is in scope, the CPI deflator extends
    # to 2025 as an ADDITIVE series (existing series id/window unchanged); the
    # extension fails loudly if the recorded CSV lacks 2025 — re-fetch it.
    if wb25 is not None:
        macro_series_series.append(cpi.extract_cpi(file_cpi, window_end=2025))
    macro_series = MacroSeriesFragment(series=macro_series_series)

    national_by_id = {s.series_id: s for s in fragment.series}
    if wb25 is not None:
        # ticket #13: the counterfactual recomputes from the latest official
        # workbook (revised 2024 + preliminary 2025). The pre-registered headline
        # window (2020-2024) stays guarded; its VALUE recomputes the same way
        # (revision policy: later official workbook wins), so the headline
        # equals the sum of the fragment's own 2020-2024 rows.
        missing_billions = compute_missing_billions(
            receipts=national_by_id["inbound_consumption_tourist_2015_2025"],
            arrivals=national_by_id["arrivals_visitor_2019_2025"],
            cpi=next(s for s in macro_series_series if s.series_id == "cpi_national_overall_2015_2025"),
            excursionist_arrivals=national_by_id["arrivals_excursionist_2019_2025"],
            land_mode_share_2024_pct=LAND_MODE_SHARE_2024_PCT,
        )
    else:
        missing_billions = compute_missing_billions(
            receipts=national_by_id["inbound_consumption_tourist_2015_2024"],
            arrivals=national_by_id["arrivals_visitor_2019_2024"],
            cpi=cpi_series,
            excursionist_arrivals=national_by_id["arrivals_excursionist_2019_2024"],
            land_mode_share_2024_pct=LAND_MODE_SHARE_2024_PCT,
        )

    # ticket T7: simulator coefficients from the source-market yields + the
    # Missing Billions deflator (additive fragment; never re-extracts anything).
    from .simulator import build_simulator_fragment

    sim_fragment = build_simulator_fragment(source_market, missing_billions)

    # ticket T5: source-market segmentation (additive fragment). The WEF TTDI
    # file defaults to the repo copy; offline tests pass a fixture instead.
    from .segmentation import build_segmentation_fragment

    wef_path = wef_csv if wef_csv is not None else _DEFAULT_WEF_CSV
    if wef_path.exists():
        segmentation = build_segmentation_fragment(source_market, wef_csv=wef_path)

    # ticket T6: regional yield benchmark (additive fragment; researched and
    # fully sourced constants, see research/regional-yield-benchmark-2024.md).
    from .regional_benchmark import build_regional_benchmark_fragment

    regional_benchmark = build_regional_benchmark_fragment()

    bundle = Bundle(
        bundle_version=BUNDLE_VERSION,
        generated_utc=datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        sources={
            file_2023.name: file_sha256(file_2023),
            file_2024.name: file_sha256(file_2024),
            file_inbrief.name: file_sha256(file_inbrief),
            file_cpi.name: file_sha256(file_cpi),
            **({file_2025.name: file_sha256(file_2025)} if wb25 is not None else {}),
            **({wef_path.name: file_sha256(wef_path)} if wef_path.exists() else {}),
        },
        fragments={
            "national_series": fragment,
            "source_market": source_market,
            "macro_series": macro_series,
            "missing_billions": missing_billions,
            "simulator": sim_fragment,
            "regional_benchmark": regional_benchmark,
            **({"source_segmentation": segmentation} if wef_path.exists() else {}),
        },
    )
    from .validate import check_cross_file_consistency, check_ground_truths

    from .validate import (
        check_cross_file_consistency,
        check_ground_truths,
        check_market_ground_truths,
        check_missing_billions,
    )

    bundle._ground_truth_report = check_ground_truths(bundle)  # type: ignore[attr-defined]
    from .validate import check_cpi_ground_truths

    bundle._cpi_report = check_cpi_ground_truths(bundle)  # type: ignore[attr-defined]
    if wb25 is not None:
        from .validate import check_2025_edition_consistency

        current_2024_series = {
            s.series_id: s
            for s in fragment.series
            if s.series_id in {
                "arrivals_visitor_2019_2024", "arrivals_tourist_2019_2024",
                "arrivals_excursionist_2019_2024", "inbound_consumption_tourist_2015_2024",
            }
        }
        bundle._edition_2025_report = check_2025_edition_consistency(
            bundle, current_2024_series
        )  # type: ignore[attr-defined]
    bundle._cross_file_report = check_cross_file_consistency(bundle, legacy)  # type: ignore[attr-defined]
    bundle._market_ground_truth_report = check_market_ground_truths(bundle)  # type: ignore[attr-defined]
    bundle._missing_billions_report = check_missing_billions(bundle)  # type: ignore[attr-defined]
    if "source_segmentation" in bundle.fragments:
        from .validate import check_segmentation_reconciliation

        bundle._segmentation_report = check_segmentation_reconciliation(bundle)  # type: ignore[attr-defined]

    # ticket T7: the simulator must reconcile with the Missing Billions headline
    # at the 2024 actual mix — checked at emission, on every run.
    from .simulator import check_simulator_reconciliation

    bundle._simulator_report = check_simulator_reconciliation(bundle)  # type: ignore[attr-defined]

    # ticket T6: the benchmark's Malaysia row must match the source-market totals
    from .validate import check_regional_benchmark

    bundle._regional_report = check_regional_benchmark(bundle)  # type: ignore[attr-defined]
    return bundle


def write_bundle(bundle: Bundle, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(bundle.to_json() + "\n")
